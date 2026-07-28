from __future__ import annotations

import json
import math
import re
from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_ALIGN_VERTICAL, WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


WORKSPACE = Path(__file__).resolve().parents[1]
PROFILE_PATH = WORKSPACE / ".database-summary-work" / "database-profile.json"
OUTPUT_PATH = WORKSPACE / "Setu_Odisha_Comprehensive_Database_Summary.docx"

with PROFILE_PATH.open("r", encoding="utf-8") as handle:
    profile = json.load(handle)


# compact_reference_guide preset, resolved to concrete values.
PAGE_WIDTH_IN = 8.5
PAGE_HEIGHT_IN = 11.0
MARGIN_IN = 1.0
HEADER_DISTANCE_IN = 0.492
FOOTER_DISTANCE_IN = 0.492
CONTENT_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120
CELL_MARGINS_DXA = {"top": 80, "bottom": 80, "start": 120, "end": 120}
FONT = "Calibri"
BODY_SIZE = 11
BODY_AFTER = 6
BODY_LINE = 1.25
H1 = {"size": 16, "color": "2E74B5", "before": 18, "after": 10}
H2 = {"size": 13, "color": "2E74B5", "before": 14, "after": 7}
H3 = {"size": 12, "color": "1F4D78", "before": 10, "after": 5}
NAVY = "0B2545"
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
MUTED = "5E6B78"
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F2F4F7"
CALLOUT = "F4F6F9"
WHITE = "FFFFFF"
BLACK = "111827"
RISK = "9B1C1C"
CAUTION = "7A5A00"


def rgb(hex_value: str) -> RGBColor:
    return RGBColor.from_string(hex_value)


def set_run_font(run, size: float | None = None, color: str | None = None,
                 bold: bool | None = None, italic: bool | None = None,
                 name: str = FONT):
    run.font.name = name
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = rgb(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_cell_shading(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")


def set_cell_margins(cell, margins=CELL_MARGINS_DXA):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in margins.items():
        tag = "w:start" if edge == "start" else "w:end" if edge == "end" else f"w:{edge}"
        node = tc_mar.find(qn(tag))
        if node is None:
            node = OxmlElement(tag)
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_table_borders(table, color="B8C2CC", size=4):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), str(size))
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), color)


def set_table_geometry(table, widths_dxa: list[int], indent_dxa: int = TABLE_INDENT_DXA):
    total = sum(widths_dxa)
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    layout = tbl_pr.first_child_found_in("w:tblLayout")
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.first_child_found_in("w:tblInd")
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for index, cell in enumerate(row.cells):
            width = widths_dxa[min(index, len(widths_dxa) - 1)]
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.first_child_found_in("w:tcW")
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            cell.width = Inches(width / 1440)
            set_cell_margins(cell)


def add_table(doc: Document, headers: list[str], rows: list[list[str]],
              widths_dxa: list[int], font_size: float = 8.5,
              header_fill: str = LIGHT_BLUE,
              alignments: list[int] | None = None,
              indent_dxa: int = TABLE_INDENT_DXA):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    header_row = table.rows[0]
    prevent_row_split(header_row)
    set_repeat_table_header(header_row)
    for idx, text in enumerate(headers):
        cell = header_row.cells[idx]
        set_cell_shading(cell, header_fill)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.keep_with_next = True
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(text)
        set_run_font(run, size=font_size, color=NAVY, bold=True)

    for row_values in rows:
        row = table.add_row()
        prevent_row_split(row)
        for idx, value in enumerate(row_values):
            cell = row.cells[idx]
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            if alignments:
                p.alignment = alignments[idx]
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            run = p.add_run(str(value))
            set_run_font(run, size=font_size, color=BLACK)
    # Named layout override: keep compact lookup/statistics tables together so a
    # one-row continuation cannot lose its context on the following page.
    if len(rows) <= 7:
        for row in table.rows[:-1]:
            for cell in row.cells:
                for p in cell.paragraphs:
                    p.paragraph_format.keep_with_next = True
    set_table_geometry(table, widths_dxa, indent_dxa)
    set_table_borders(table)
    after = doc.add_paragraph()
    after.paragraph_format.space_before = Pt(0)
    after.paragraph_format.space_after = Pt(4)
    after.paragraph_format.line_spacing = 1
    return table


def add_page_field(paragraph):
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char1, instr, fld_char2])
    set_run_font(run, size=8.5, color=MUTED)


def configure_document(doc: Document):
    section = doc.sections[0]
    section.page_width = Inches(PAGE_WIDTH_IN)
    section.page_height = Inches(PAGE_HEIGHT_IN)
    section.top_margin = Inches(MARGIN_IN)
    section.bottom_margin = Inches(MARGIN_IN)
    section.left_margin = Inches(MARGIN_IN)
    section.right_margin = Inches(MARGIN_IN)
    section.header_distance = Inches(HEADER_DISTANCE_IN)
    section.footer_distance = Inches(FOOTER_DISTANCE_IN)

    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    normal.font.size = Pt(BODY_SIZE)
    normal.font.color.rgb = rgb(BLACK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(BODY_AFTER)
    normal.paragraph_format.line_spacing = BODY_LINE

    for style_name, token in (("Heading 1", H1), ("Heading 2", H2), ("Heading 3", H3)):
        style = doc.styles[style_name]
        style.font.name = FONT
        style._element.rPr.rFonts.set(qn("w:ascii"), FONT)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
        style.font.size = Pt(token["size"])
        style.font.color.rgb = rgb(token["color"])
        style.font.bold = True
        style.paragraph_format.space_before = Pt(token["before"])
        style.paragraph_format.space_after = Pt(token["after"])
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.keep_together = True

    for style_name in ("List Bullet", "List Number"):
        style = doc.styles[style_name]
        style.font.name = FONT
        style._element.rPr.rFonts.set(qn("w:ascii"), FONT)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
        style.font.size = Pt(BODY_SIZE)
        style.paragraph_format.left_indent = Inches(0.375)
        style.paragraph_format.first_line_indent = Inches(-0.188)
        style.paragraph_format.space_before = Pt(0)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = BODY_LINE

    for sec in doc.sections:
        header = sec.header
        p = header.paragraphs[0]
        p.text = ""
        p.paragraph_format.space_after = Pt(0)
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        run = p.add_run("SETU_ODISHA  |  DATABASE TECHNICAL SUMMARY")
        set_run_font(run, size=8.5, color=MUTED, bold=True)
        footer = sec.footer
        fp = footer.paragraphs[0]
        fp.text = ""
        fp.paragraph_format.space_before = Pt(0)
        fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        label = fp.add_run("Page ")
        set_run_font(label, size=8.5, color=MUTED)
        add_page_field(fp)


def add_heading(doc, text: str, level: int):
    p = doc.add_heading(text, level=level)
    p.paragraph_format.keep_with_next = True
    return p


def add_body(doc, text: str, bold_label: str | None = None):
    p = doc.add_paragraph()
    if bold_label:
        r = p.add_run(bold_label)
        set_run_font(r, bold=True, color=NAVY)
    r = p.add_run(text)
    set_run_font(r)
    return p


def add_bullet(doc, text: str, color: str | None = None):
    p = doc.add_paragraph(style="List Bullet")
    r = p.add_run(text)
    set_run_font(r, color=color or BLACK)
    return p


def add_callout(doc, label: str, text: str, fill: str = CALLOUT):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.08)
    p.paragraph_format.right_indent = Inches(0.08)
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(9)
    p.paragraph_format.line_spacing = 1.15
    p_pr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")
    p_pr.append(shd)
    r = p.add_run(f"{label}: ")
    set_run_font(r, size=10.5, bold=True, color=NAVY)
    r = p.add_run(text)
    set_run_font(r, size=10.5, color=BLACK)


def clean_text(value) -> str:
    if value is None:
        return ""
    text = str(value).replace("\r", " ").replace("\n", " ").replace("\t", " ")
    return re.sub(r"\s+", " ", text).strip()


def humanize(name: str) -> str:
    value = name.replace("_", " ")
    value = re.sub(r"([a-z0-9])([A-Z])", r"\1 \2", value)
    value = re.sub(r"([A-Z]+)([A-Z][a-z])", r"\1 \2", value)
    return re.sub(r"\s+", " ", value).strip()


def exact_rows(table_name: str) -> int:
    item = next(x for x in profile["profiles"] if x["table_name"] == table_name)
    return int(item.get("exact_row_count", 0))


def table_category(name: str) -> str:
    if name.startswith("AspNet"):
        return "Security / identity"
    if name in {"__EFMigrationsHistory", "sysdiagrams"}:
        return "Platform / administration"
    if name.startswith("Deleted"):
        return "Archive / deletion audit"
    if name in {"PatientTable", "TumourTable", "SourceTable", "PatientFile", "PossibleDuplicates"}:
        return "Core / transactional"
    if name in {"Centres", "StaffDetails"}:
        return "Operational master"
    if name.endswith("List") or name in {"ICCCMaster", "ICD10Group", "Lymph_Leuk", "Meso_Mela", "Unknown"}:
        return "Reference / classification"
    return "Supporting"


PURPOSES = {
    "__EFMigrationsHistory": "Records applied Entity Framework Core migrations for application schema versioning.",
    "PatientTable": "Stores the active patient-level cancer registry record and demographic/vital-status attributes.",
    "TumourTable": "Stores tumour/case attributes linked logically to a patient, including diagnosis, site, histology and treatment coding.",
    "SourceTable": "Stores source-of-information records linked logically to a tumour and patient.",
    "PatientFile": "Stores file metadata for documents attached logically to tumour records; file contents are not stored in this table.",
    "PossibleDuplicates": "Tracks suspected duplicate record pairs, match scores, decisions and review workflow metadata.",
    "DeletedPatientTable": "Archives deleted patient rows together with deletion and audit metadata.",
    "DeletedTumourTable": "Archives deleted tumour rows together with deletion and audit metadata.",
    "DeletedSourceTable": "Archives deleted source rows together with deletion and audit metadata.",
    "Centres": "Master list of participating centres and their geographic codes.",
    "StaffDetails": "Intended to store staff master and contact/assignment metadata; currently empty.",
    "StateList": "State-level geographic master.",
    "DistrictList": "District-level geographic master keyed to state codes.",
    "BlockList": "Block-level geographic master keyed to state and district codes.",
    "VillageList": "Village-level geographic master keyed to state, district and block codes.",
    "ICD10Group": "Maps normalized ICD-10 codes to grouped cancer-site labels used by incidence and mortality queries.",
    "ICCCMaster": "Intended International Classification of Childhood Cancer reference structure; currently empty.",
    "Lymph_Leuk": "Intended lymphoma/leukaemia classification lookup; currently empty.",
    "Meso_Mela": "Intended mesothelioma/melanoma classification lookup; currently empty.",
    "sysdiagrams": "SQL Server Management Studio database-diagram metadata table.",
    "Unknown": "Generic code/value placeholder reference table; currently empty.",
}


def table_purpose(name: str) -> tuple[str, str]:
    if name in PURPOSES:
        confidence = "Confirmed by schema/application use" if name in {
            "__EFMigrationsHistory", "PatientTable", "TumourTable", "SourceTable",
            "PatientFile", "PossibleDuplicates", "ICD10Group", "sysdiagrams"
        } else "Purpose inferred from table and column names"
        return PURPOSES[name], confidence
    if name.startswith("AspNet"):
        suffix = humanize(name.removeprefix("AspNet"))
        return (
            f"Standard ASP.NET Core Identity table for {suffix.lower()} data.",
            "Confirmed standard framework schema",
        )
    if name.endswith("List"):
        subject = humanize(name[:-4]).lower()
        return (
            f"Code/value or geographic reference data for {subject}.",
            "Purpose inferred from table and column names",
        )
    return (
        f"Supporting table for {humanize(name).lower()} data.",
        "Purpose inferred from table and column names",
    )


COLUMN_DESCRIPTIONS = {
    "Id": "Internal row identifier.",
    "Code": "Stored category or classification code.",
    "Value": "Display label or decoded meaning for the associated code.",
    "REGNO": "Registry number used as the logical patient link.",
    "TUMOURID": "Logical tumour identifier.",
    "SOURCEID": "Logical source-record identifier.",
    "PATIENTIDTUMOURTABLE": "Patient registry key used by the application to link tumour rows to PatientTable.REGNO.",
    "TUMOURIDSOURCETABLE": "Tumour key used to link source rows to TumourTable.TUMOURID.",
    "CentreCode": "Code identifying the participating registry centre.",
    "StateId": "State master code.",
    "DistrictId": "District master code.",
    "BlockId": "Block master code.",
    "VillageId": "Village master code.",
    "ICD10": "ICD-10 cancer diagnosis/site code.",
    "CreatedDate": "Record creation timestamp.",
    "LastModifiedDate": "Latest recorded modification timestamp.",
    "DeletedDate": "Deletion/archive timestamp.",
    "UploadedDate": "File metadata upload timestamp.",
    "FileSize": "Recorded file size in bytes.",
    "ContentType": "MIME media type for the referenced file.",
    "MigrationId": "Entity Framework migration identifier.",
    "ProductVersion": "Entity Framework product version used for the migration.",
}


def data_type_text(column: dict) -> str:
    data_type = column["data_type"]
    max_length = column.get("max_length")
    precision = column.get("precision")
    scale = column.get("scale")
    if data_type in {"nvarchar", "varchar", "nchar", "char", "binary", "varbinary"}:
        length = "max" if max_length == -1 else str(max_length)
        return f"{data_type}({length})"
    if data_type in {"decimal", "numeric"}:
        return f"{data_type}({precision},{scale})"
    if data_type in {"datetime2", "datetimeoffset", "time"}:
        return f"{data_type}({scale})"
    return data_type


def get_column_profile(table_name: str, column_name: str) -> dict:
    table_profile = next(x for x in profile["profiles"] if x["table_name"] == table_name)
    return next(x for x in table_profile["columns"] if x["column_name"] == column_name)


def representative_text(column_profile: dict) -> str:
    if column_profile.get("sensitive"):
        return "Representative values masked (sensitive or identifying field)."
    values = column_profile.get("representative_values")
    if not isinstance(values, list) or not values:
        return ""
    rendered = []
    for item in values[:3]:
        value = clean_text(item.get("value"))
        if not value:
            continue
        if len(value) > 55:
            value = value[:52] + "..."
        rendered.append(f"{value} ({item.get('frequency', 0):,})")
    return "Representative values: " + "; ".join(rendered) + "." if rendered else ""


def column_description(table_name: str, column: dict) -> str:
    name = column["column_name"]
    if column.get("description"):
        base = f"Catalog description: {clean_text(column['description'])}"
    elif name in COLUMN_DESCRIPTIONS:
        base = COLUMN_DESCRIPTIONS[name]
    else:
        base = f"Inferred: stores {humanize(name).lower()} data."
    cp = get_column_profile(table_name, name)
    row_count = exact_rows(table_name)
    facts = []
    if row_count:
        null_count = cp.get("null_count")
        blank_count = cp.get("blank_count")
        missing = (int(null_count or 0) + int(blank_count or 0))
        if missing:
            facts.append(f"{missing:,}/{row_count:,} rows are null or blank")
        distinct = cp.get("distinct_count")
        if distinct is not None and not cp.get("sensitive"):
            facts.append(f"{int(distinct):,} distinct non-null values")
        if cp.get("min_value") is not None or cp.get("max_value") is not None:
            facts.append(
                f"range {clean_text(cp.get('min_value')) or 'n/a'} to "
                f"{clean_text(cp.get('max_value')) or 'n/a'}"
            )
    if facts:
        base += " Profile: " + "; ".join(facts) + "."
    sample = representative_text(cp)
    if sample:
        base += " " + sample
    return base


def constraint_text(column: dict) -> str:
    parts = []
    if column.get("is_primary_key"):
        ordinal = column.get("primary_key_ordinal")
        parts.append(f"PK#{ordinal}" if ordinal else "PK")
    if column.get("is_identity"):
        seed = clean_text(column.get("seed_value"))
        inc = clean_text(column.get("increment_value"))
        parts.append(f"IDENTITY({seed},{inc})" if seed and inc else "IDENTITY")
    if column.get("is_foreign_key"):
        parts.append(
            f"FK -> {column.get('referenced_schema')}.{column.get('referenced_table')}"
            f".{column.get('referenced_column')}"
        )
    if column.get("is_unique_key"):
        parts.append("UNIQUE INDEX")
    if column.get("is_computed"):
        parts.append("COMPUTED")
    return "; ".join(parts) or "None"


def table_relationships(table_name: str) -> list[str]:
    relationships = []
    for fk in profile["foreign_keys"]:
        if fk["child_table"] == table_name:
            relationships.append(
                f"Enforced many-to-one: dbo.{table_name}.{fk['child_column']} -> "
                f"{fk['parent_schema']}.{fk['parent_table']}.{fk['parent_column']} "
                f"(ON DELETE {fk['on_delete']}; ON UPDATE {fk['on_update']})."
            )
        if fk["parent_table"] == table_name:
            relationships.append(
                f"Enforced one-to-many parent for {fk['child_schema']}.{fk['child_table']} "
                f"through {fk['foreign_key_name']}."
            )
    logical = {
        "PatientTable": [
            "Validated logical parent of TumourTable through PatientTable.REGNO = TumourTable.PATIENTIDTUMOURTABLE; 1,455/1,455 tumour rows matched.",
            "Validated logical parent of SourceTable through REGNO; 1,455/1,455 source rows matched.",
        ],
        "TumourTable": [
            "Validated logical child of PatientTable and parent of SourceTable.",
            "Validated logical parent of PatientFile through TumourTable.TUMOURID = PatientFile.TumourId; 1,349/1,349 file rows matched.",
        ],
        "SourceTable": [
            "Validated logical child of TumourTable through TUMOURIDSOURCETABLE and of PatientTable through REGNO; all current rows matched.",
        ],
        "PatientFile": [
            "Validated logical child of TumourTable through TumourId; all 1,349 current rows matched.",
        ],
        "StateList": ["Logical one-to-many geographic parent of DistrictList; all current district state codes matched."],
        "DistrictList": ["Logical child of StateList and parent of BlockList; all current block district codes matched."],
        "BlockList": ["Logical child of DistrictList and parent of VillageList; 310 VillageList rows reference a BlockId absent from BlockList."],
        "VillageList": ["Logical child of BlockList and geographic parent of Centres; all current centre village codes matched."],
        "Centres": ["Logically references StateList, DistrictList, BlockList and VillageList by stored geographic codes; no database FKs enforce these links."],
        "ICD10Group": ["Application query lookup for normalized TumourTable.ICD10 values; 11 populated tumour codes were unmapped and 2 failed the expected Cnn pattern."],
    }
    relationships.extend(logical.get(table_name, []))

    semantic = {
        "AreaList": "PatientTable.Area",
        "BasisOfDiagnosisList": "TumourTable.BasisOfDiagnosis",
        "BehaviourList": "TumourTable.Behaviour",
        "CauseOfDeathList": "PatientTable.CauseOfDeath",
        "CentreTypeList": "Centres.CentreType",
        "ClinicalStageList": "TumourTable.ClinicalStage",
        "GradeList": "TumourTable.GradeCellType",
        "HistologyList": "TumourTable.Histology",
        "PrimarySiteList": "TumourTable.PrimarySite",
        "SexList": "PatientTable.Sex",
        "TreatmentList": "TumourTable.Treatment",
        "TreatmentStatusList": "TumourTable.TreatmentStatus",
        "TypeOfSourceList": "SourceTable.TypeOfSource",
        "VitalStatusList": "PatientTable.VitalStatus",
    }
    if table_name in semantic:
        relationships.append(
            f"Probable semantic code/value lookup for {semantic[table_name]}, based on matching domain names; not enforced by a foreign key."
        )
    return relationships or ["No enforced or validated logical relationships were identified for this table."]


def table_quality_observations(table_name: str) -> list[str]:
    rows = exact_rows(table_name)
    columns = [x for x in profile["columns"] if x["table_name"] == table_name]
    table_profile = next(x for x in profile["profiles"] if x["table_name"] == table_name)
    observations = []
    if rows == 0:
        observations.append("The table is empty at the profiling timestamp; value distributions and duplicates cannot be assessed.")
    else:
        missing = []
        for cp in table_profile["columns"]:
            count = int(cp.get("null_count") or 0) + int(cp.get("blank_count") or 0)
            if count:
                missing.append((count, cp["column_name"]))
        if missing:
            missing.sort(reverse=True)
            top = ", ".join(f"{name} {count:,}/{rows:,}" for count, name in missing[:5])
            observations.append(f"Highest null/blank counts: {top}.")
        else:
            observations.append("No null or blank values were found in profiled columns.")

    pk_columns = [c for c in columns if c.get("is_primary_key")]
    if pk_columns:
        observations.append("The primary key prevents duplicate full rows and duplicate key values.")
    else:
        code_profile = next((x for x in table_profile["columns"] if x["column_name"] == "Code"), None)
        if rows and code_profile and code_profile.get("distinct_count") == rows:
            observations.append("Code values are unique in the current data, but no PK/UNIQUE constraint enforces that condition.")
        elif rows:
            observations.append("No primary key exists; duplicate row prevention is not enforced.")

    special = {
        "PatientTable": [
            "REGNO has 0 duplicate groups; all 1,450 values are represented in logical tumour/source links as expected.",
            "DateOfBirth and DateOfDeath are nvarchar(max), but all populated current values converted to date in the read-only check.",
            "ContactNumber is stored as float, which is not a safe type for exact telephone identifiers.",
        ],
        "TumourTable": [
            "TUMOURID has 0 duplicate groups and all patient links match, despite the absence of an FK/UNIQUE constraint.",
            "ICD10 contains 2 values outside the expected normalized Cnn pattern and 11 values with no ICD10Group match.",
            "DateOfDiagnosis is nvarchar(max); all 1,412 populated values converted to date in the read-only check.",
        ],
        "SourceTable": [
            "SOURCEID has 0 duplicate groups; all patient and tumour logical links match.",
        ],
        "PatientFile": [
            "All 1,349 TumourId values matched a TumourTable.TUMOURID.",
        ],
        "VillageList": [
            "VillageId has 66 duplicate groups (66 excess rows) and is not protected by a unique constraint.",
            "310 village rows reference a BlockId absent from BlockList.",
        ],
        "PossibleDuplicates": [
            "All 37 rows contain a decision, but ReviewedBy is null/blank in all 37 rows; review accountability is incomplete.",
        ],
        "AspNetUsers": [
            "154 users exist while the standard role, role-claim, user-role, user-claim, login and token tables are empty.",
            "Authorization appears to rely on the custom UserType/CentreCode fields or external Windows authentication; verify intended policy.",
        ],
        "ICD10Group": [
            "Code is unique across 87 rows; Value has 59 distinct labels, consistent with multiple codes mapping to common grouped sites.",
        ],
    }
    observations.extend(special.get(table_name, []))
    if any(c["data_type"] == "nvarchar" and c.get("max_length") == -1 for c in columns):
        count = sum(1 for c in columns if c["data_type"] == "nvarchar" and c.get("max_length") == -1)
        observations.append(f"{count} column(s) use nvarchar(max), limiting type precision and potentially increasing storage/indexing cost.")
    return observations


def important_columns(table_name: str) -> list[str]:
    columns = [x for x in profile["columns"] if x["table_name"] == table_name]
    selected = [
        c["column_name"] for c in columns
        if c.get("is_primary_key") or c.get("is_foreign_key")
        or re.search(r"(REGNO|TUMOURID|SOURCEID|Code$|Status|Date|Id$)", c["column_name"], re.I)
    ]
    return selected[:12]


doc = Document()
configure_document(doc)

# Editorial-cover opening pattern.
for _ in range(4):
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(18)

kicker = doc.add_paragraph()
kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
kicker.paragraph_format.space_after = Pt(18)
run = kicker.add_run("TECHNICAL DATABASE REPORT")
set_run_font(run, size=11, color=BLUE, bold=True)

title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.paragraph_format.space_after = Pt(8)
run = title.add_run("Comprehensive Database Summary")
set_run_font(run, size=30, color=NAVY, bold=True)

subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
subtitle.paragraph_format.space_after = Pt(4)
run = subtitle.add_run("Setu_Odisha")
set_run_font(run, size=18, color=DARK_BLUE, bold=True)

scope = doc.add_paragraph()
scope.alignment = WD_ALIGN_PARAGRAPH.CENTER
scope.paragraph_format.space_after = Pt(26)
run = scope.add_run("Schema, relationships, database objects, record counts, column dictionary and data-quality analysis")
set_run_font(run, size=11, color=MUTED, italic=True)

generated_at = datetime.fromisoformat(profile["generated_at"]).astimezone()
meta = doc.add_paragraph()
meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
meta.paragraph_format.space_after = Pt(4)
run = meta.add_run(f"Profile generated {generated_at:%d %B %Y, %H:%M %Z}")
set_run_font(run, size=10, color=MUTED)

meta = doc.add_paragraph()
meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = meta.add_run("Read-only catalog and aggregate-data analysis; sensitive values masked")
set_run_font(run, size=10, color=MUTED)

doc.add_page_break()

add_heading(doc, "Detailed Summary of Overall Database", 1)

add_heading(doc, "Database Overview", 2)
db = profile["database"]
add_body(
    doc,
    f"Setu_Odisha is an online, multi-user Microsoft SQL Server database hosted on "
    f"{db['server_name']}. It was created on {clean_text(db['create_date'])}, uses compatibility "
    f"level {db['compatibility_level']}, collation {db['collation_name']}, and the "
    f"{db['recovery_model_desc']} recovery model. The database is writable and has "
    f"{float(db['allocated_size_mb']):,.2f} MB allocated.",
)
add_body(
    doc,
    "Purpose (confirmed from the application query paths and schema): it supports the Odisha population-based cancer registry portal, "
    "including patient, tumour, source, file, duplicate-review, geographic, classification and user-account data.",
)
add_callout(
    doc,
    "Evidence boundary",
    "Catalog facts, exact record counts and aggregate checks are confirmed directly from the database. "
    "Business descriptions without database extended properties are explicitly labeled as inferred from names, columns or application use. "
    "No patient/user-identifiable sample values are reproduced.",
)

add_heading(doc, "Database Statistics", 2)
total_rows = sum(int(x.get("exact_row_count", 0)) for x in profile["profiles"])
pk_tables = len({x["table_name"] for x in profile["columns"] if x.get("is_primary_key")})
empty_tables = sum(1 for x in profile["profiles"] if int(x.get("exact_row_count", 0)) == 0)
type_counts = Counter(x["data_type"] for x in profile["columns"])
object_counts = Counter(x["type_desc"] for x in profile["objects"])
stats_rows = [
    ["Database", db["database_name"]],
    ["Base tables", f"{len(profile['tables']):,}"],
    ["Total exact records", f"{total_rows:,}"],
    ["Total columns", f"{len(profile['columns']):,}"],
    ["Populated / empty tables", f"{len(profile['tables']) - empty_tables:,} / {empty_tables:,}"],
    ["Tables with a primary key", f"{pk_tables:,}"],
    ["Indexes", f"{len(profile['indexes']):,} (including PK and unique indexes)"],
    ["Enforced foreign keys", f"{len(profile['foreign_keys']):,}"],
    ["CHECK / DEFAULT constraints", f"{len(profile['checks']):,} / {sum(1 for x in profile['columns'] if x.get('default_definition')):,}"],
    ["Views / stored procedures", f"{object_counts.get('VIEW', 0):,} / {object_counts.get('SQL_STORED_PROCEDURE', 0):,}"],
    ["Functions / triggers", f"{sum(v for k, v in object_counts.items() if 'FUNCTION' in k):,} / {sum(v for k, v in object_counts.items() if 'TRIGGER' in k):,}"],
    ["Allocated size", f"{float(db['allocated_size_mb']):,.2f} MB"],
]
add_table(doc, ["Metric", "Confirmed value"], stats_rows, [2700, 6660], font_size=9.3)

type_rows = [[name, f"{count:,}", f"{count / len(profile['columns']):.1%}"] for name, count in type_counts.most_common()]
add_heading(doc, "Data-Type Distribution", 3)
add_table(doc, ["Data type", "Columns", "Share"], type_rows, [3600, 2500, 3260], font_size=9.2,
          alignments=[WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER])
add_body(
    doc,
    f"Confirmed profile: {type_counts.get('nvarchar', 0)} of {len(profile['columns'])} columns are nvarchar, "
    f"and {sum(1 for x in profile['columns'] if x['data_type'] == 'nvarchar' and x.get('max_length') == -1)} "
    "are nvarchar(max). This schema is strongly text-oriented.",
)

add_heading(doc, "Database Architecture", 2)
add_body(
    doc,
    "The database is a single-schema (dbo) relational design with four functional layers:",
)
for text in [
    "Core registry transactions: PatientTable, TumourTable, SourceTable, PatientFile and PossibleDuplicates.",
    "Reference/master data: geography, centres, cancer classifications, clinical codes and other Code/Value lookup tables.",
    "Archive/audit data: DeletedPatientTable, DeletedTumourTable and DeletedSourceTable.",
    "Platform/security data: ASP.NET Core Identity tables, Entity Framework migration history and SQL Server diagram metadata.",
]:
    add_bullet(doc, text)
add_body(
    doc,
    "Core data flow (validated): PatientTable.REGNO -> TumourTable.PATIENTIDTUMOURTABLE -> "
    "SourceTable.TUMOURIDSOURCETABLE, with PatientFile.TumourId also pointing logically to TumourTable.TUMOURID. "
    "All current rows in those four logical child links matched a parent. These relationships are not enforced by database foreign keys.",
)

add_heading(doc, "Tables and Their Roles", 2)
directory_rows = []
for table in profile["tables"]:
    purpose, _ = table_purpose(table["table_name"])
    directory_rows.append([
        f"dbo.{table['table_name']}",
        table_category(table["table_name"]),
        f"{exact_rows(table['table_name']):,}",
        purpose,
    ])
add_table(
    doc,
    ["Table", "Role", "Rows", "Purpose"],
    directory_rows,
    [2100, 1700, 900, 4660],
    font_size=8.0,
    alignments=[WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.RIGHT, WD_ALIGN_PARAGRAPH.LEFT],
)

add_heading(doc, "Relationships and Dependencies", 2)
add_heading(doc, "Enforced foreign keys", 3)
fk_rows = []
for fk in profile["foreign_keys"]:
    fk_rows.append([
        fk["foreign_key_name"],
        f"{fk['child_table']}.{fk['child_column']}",
        f"{fk['parent_table']}.{fk['parent_column']}",
        fk["on_delete"],
        "Trusted" if not fk["is_not_trusted"] and not fk["is_disabled"] else "Review",
    ])
add_table(doc, ["Constraint", "Child", "Parent", "Delete", "Status"], fk_rows,
          [2700, 1900, 1900, 1200, 1660], font_size=7.9)
add_body(
    doc,
    "Confirmed: all six enforced foreign keys belong to the ASP.NET Identity subsystem. "
    "No cancer-registry, file, centre, geography or clinical lookup relationship is enforced by an FK.",
)

add_heading(doc, "Validated logical relationships", 3)
logical_rows = []
for row in profile["targeted_checks"]["logical_link_coverage"]:
    logical_rows.append([
        row["relationship"],
        f"{int(row['child_rows']):,}",
        f"{int(row['matched_rows'] or 0):,}",
        f"{int(row['orphan_rows'] or 0):,}",
    ])
file_link = profile["targeted_checks"]["patient_file_link"][0]
logical_rows.append([
    "PatientFile.TumourId -> TumourTable.TUMOURID",
    f"{int(file_link['file_rows']):,}",
    f"{int(file_link['matched_rows']):,}",
    f"{int(file_link['orphan_rows']):,}",
])
add_table(doc, ["Logical relationship", "Child rows", "Matched", "Orphans"],
          logical_rows, [5000, 1450, 1450, 1460], font_size=8.5,
          alignments=[WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.RIGHT, WD_ALIGN_PARAGRAPH.RIGHT, WD_ALIGN_PARAGRAPH.RIGHT])

card_rows = []
for row in profile["targeted_checks"]["relationship_cardinality"]:
    card_rows.append([
        row["relationship"],
        f"{int(row['parent_rows']):,}",
        f"{int(row['parents_without_children'] or 0):,}",
        f"{int(row['parents_with_one_child'] or 0):,}",
        f"{int(row['parents_with_multiple_children'] or 0):,}",
        f"{int(row['max_children'] or 0):,}",
    ])
add_heading(doc, "Observed cardinality", 3)
add_table(doc, ["Relationship", "Parents", "No child", "One", "Many", "Max"],
          card_rows, [3600, 1150, 1150, 1100, 1150, 1210], font_size=8.2,
          alignments=[WD_ALIGN_PARAGRAPH.LEFT] + [WD_ALIGN_PARAGRAPH.RIGHT] * 5)

add_heading(doc, "Keys, Constraints, and Indexes", 2)
add_body(
    doc,
    f"The catalog contains {len(profile['indexes'])} indexes across {len({x['table_name'] for x in profile['indexes']})} tables. "
    f"{pk_tables} tables have primary keys; {len(profile['tables']) - pk_tables} do not. "
    "There are no user-defined CHECK constraints, no column DEFAULT constraints and no computed columns.",
)
index_rows = []
for idx in profile["indexes"]:
    index_rows.append([
        idx["table_name"],
        idx["index_name"],
        "PK" if idx["is_primary_key"] else "Unique" if idx["is_unique"] else "Nonunique",
        clean_text(idx.get("columns")),
    ])
add_table(doc, ["Table", "Index", "Type", "Columns"], index_rows,
          [1900, 3000, 1200, 3260], font_size=7.9)

add_heading(doc, "Views, Stored Procedures, Functions, and Triggers", 2)
add_body(
    doc,
    f"Confirmed catalog totals: 0 views; {object_counts.get('SQL_STORED_PROCEDURE', 0)} stored procedures; "
    f"{sum(v for k, v in object_counts.items() if 'FUNCTION' in k)} function; 0 triggers.",
)
parameters_by_object = defaultdict(list)
for param in profile["object_parameters"]:
    parameters_by_object[param["object_name"]].append(param)
deps_by_object = defaultdict(list)
for dep in profile["object_dependencies"]:
    deps_by_object[dep["referencing_object"]].append(
        ".".join(x for x in [dep.get("referenced_schema_name"), dep.get("referenced_entity_name")] if x)
    )
object_rows = []
for obj in profile["objects"]:
    params = ", ".join(
        f"{p['parameter_name']} {p['data_type']}"
        for p in parameters_by_object.get(obj["object_name"], [])
    ) or "None"
    deps = ", ".join(sorted(set(deps_by_object.get(obj["object_name"], [])))) or "None recorded"
    if obj["object_name"].startswith("SP_GET_"):
        purpose = "Application export/report procedure; purpose inferred from object name."
    elif "diagram" in obj["object_name"].lower():
        purpose = "SQL Server diagram-management support object."
    else:
        purpose = "Database function/procedure; inspect definition before changing."
    object_rows.append([
        obj["type_desc"].replace("SQL_", "").replace("_", " ").title(),
        f"{obj['schema_name']}.{obj['object_name']}",
        params,
        deps,
        purpose,
    ])
add_table(doc, ["Type", "Object", "Parameters", "Dependencies", "Purpose"],
          object_rows, [1200, 2300, 2100, 1700, 2060], font_size=7.4)

add_heading(doc, "Data Quality and Integrity Analysis", 2)
confirmed_findings = [
    "Core logical integrity is currently strong: all 1,455 tumour-to-patient, 1,455 source-to-tumour, 1,455 source-to-patient and 1,349 file-to-tumour links matched.",
    "Business identifiers REGNO, TUMOURID and SOURCEID have no duplicate groups in current data, but uniqueness is not enforced.",
    "VillageList contains 66 duplicate VillageId groups (66 excess rows) and 310 rows whose BlockId is absent from BlockList.",
    "TumourTable.ICD10 has 2 populated values outside the expected normalized Cnn pattern and 11 populated values not mapped by ICD10Group.",
    "PossibleDuplicates contains 37 decisions; ReviewedBy is missing in all 37 rows.",
    f"{empty_tables} of {len(profile['tables'])} tables are empty, including several classification/reference tables and all standard Identity role/claim/link tables.",
    f"{sum(1 for x in profile['columns'] if x.get('is_nullable'))} of {len(profile['columns'])} columns are nullable; "
    f"{sum(1 for x in profile['columns'] if x['data_type'] == 'nvarchar' and x.get('max_length') == -1)} are nvarchar(max).",
    "No user-defined CHECK constraints or DEFAULT constraints exist.",
    "All populated DateOfBirth, DateOfDeath and DateOfDiagnosis values converted successfully in the current read-only check; the fields remain stored as text.",
]
for item in confirmed_findings:
    add_bullet(doc, item)

add_heading(doc, "Important Findings and Recommendations", 2)
recommendations = [
    ("High", "Repair the 310 VillageList rows with unknown BlockId and resolve the 66 duplicate VillageId groups before introducing constraints."),
    ("High", "Add staged UNIQUE constraints for PatientTable.REGNO, TumourTable.TUMOURID and SourceTable.SOURCEID after regression testing."),
    ("High", "Add staged FKs for the validated patient-tumour-source-file chain and geographic hierarchy; use WITH CHECK so trust is verified."),
    ("High", "Remediate the 2 malformed and 11 unmapped ICD10 values, then enforce normalized code validation or a controlled lookup."),
    ("Medium", "Migrate date strings to date/datetime2 columns and ContactNumber away from float; preserve source text during a controlled transition."),
    ("Medium", "Replace unnecessary nvarchar(max) columns with bounded lengths derived from real maxima, prioritizing keys and frequently filtered fields."),
    ("Medium", "Add CHECK/default rules for status flags, required workflow fields and audit timestamps; define nullability from business rules."),
    ("Medium", "Require reviewer identity for duplicate decisions or explicitly record an automated decision source."),
    ("Medium", "Review the 25 empty tables and retire, populate or document them so inactive schema does not obscure intended functionality."),
    ("Security", "Verify the intended authorization model: 154 AspNetUsers rows exist while role/claim membership tables are empty and the web app also uses Windows authentication."),
]
rec_rows = [[severity, recommendation] for severity, recommendation in recommendations]
add_table(doc, ["Priority", "Recommendation"], rec_rows, [1400, 7960], font_size=9.0,
          alignments=[WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT])

doc.add_page_break()
add_heading(doc, "Detailed Summary of Individual Tables", 1)
add_body(
    doc,
    "Each section below reports the exact record count, catalog metadata, every column, enforced and logical relationships, "
    "aggregate profile facts and privacy-safe representative values. Descriptions prefixed “Inferred” are observations, not catalog documentation.",
)

tables_sorted = profile["tables"]
columns_by_table = defaultdict(list)
for column in profile["columns"]:
    columns_by_table[column["table_name"]].append(column)

for number, table in enumerate(tables_sorted, start=1):
    table_name = table["table_name"]
    rows = exact_rows(table_name)
    columns = columns_by_table[table_name]
    purpose, confidence = table_purpose(table_name)
    heading = add_heading(doc, f"{number}. dbo.{table_name}", 2)
    heading.paragraph_format.page_break_before = False

    add_heading(doc, "Table Overview", 3)
    add_body(doc, purpose + " ", bold_label=f"{confidence}. ")
    add_body(doc, f"Role classification: {table_category(table_name)}.")

    add_heading(doc, "Table Statistics", 3)
    pk_names = [c["column_name"] for c in columns if c.get("is_primary_key")]
    fk_count = sum(1 for c in columns if c.get("is_foreign_key"))
    stats = [
        ["Total records", f"{rows:,}"],
        ["Columns", f"{len(columns):,}"],
        ["Primary key", ", ".join(pk_names) if pk_names else "None"],
        ["Enforced foreign-key columns", f"{fk_count:,}"],
        ["Created", clean_text(table.get("create_date"))],
        ["Last schema modification", clean_text(table.get("modify_date"))],
    ]
    add_table(doc, ["Measure", "Value"], stats, [2500, 6860], font_size=8.8)

    add_heading(doc, "Column Details", 3)
    column_rows = []
    for column in columns:
        column_rows.append([
            column["column_name"],
            data_type_text(column),
            "Yes" if column.get("is_nullable") else "No",
            constraint_text(column),
            clean_text(column.get("default_definition")) or "None",
            column_description(table_name, column),
        ])
    add_table(
        doc,
        ["Column Name", "Data Type", "Nullable", "Key / Constraint", "Default", "Description"],
        column_rows,
        [1450, 1200, 850, 1650, 950, 3260],
        font_size=7.45,
        alignments=[
            WD_ALIGN_PARAGRAPH.LEFT,
            WD_ALIGN_PARAGRAPH.LEFT,
            WD_ALIGN_PARAGRAPH.CENTER,
            WD_ALIGN_PARAGRAPH.LEFT,
            WD_ALIGN_PARAGRAPH.LEFT,
            WD_ALIGN_PARAGRAPH.LEFT,
        ],
    )

    add_heading(doc, "Relationships", 3)
    for relationship in table_relationships(table_name):
        add_bullet(doc, relationship)

    add_heading(doc, "Data Analysis", 3)
    important = important_columns(table_name)
    add_body(
        doc,
        "Business-critical or frequently joined/filterable columns: "
        + (", ".join(important) if important else "No specific critical column was identifiable from keys and names.")
        + ".",
    )
    for observation in table_quality_observations(table_name):
        add_bullet(doc, observation)

    add_heading(doc, "Data Quality Observations", 3)
    if rows == 0:
        add_body(doc, "Confirmed: no rows are present; completeness, consistency and representative-value checks are not applicable.")
    else:
        tp = next(x for x in profile["profiles"] if x["table_name"] == table_name)
        complete_columns = 0
        for cp in tp["columns"]:
            missing = int(cp.get("null_count") or 0) + int(cp.get("blank_count") or 0)
            if missing == 0:
                complete_columns += 1
        add_body(
            doc,
            f"{complete_columns} of {len(columns)} columns have no null/blank values in the current data. "
            "Representative values in the column table are frequency-based and masked for sensitive fields.",
        )

    add_heading(doc, "Important Findings", 3)
    findings = []
    if not pk_names:
        findings.append("No primary key is defined.")
    if fk_count == 0 and any(x for x in table_relationships(table_name) if "logical" in x.lower() or "semantic" in x.lower()):
        findings.append("Application-level relationships exist without FK enforcement.")
    if rows == 0:
        findings.append("The table is inactive/empty at the profiling timestamp; confirm whether it is planned, legacy or awaiting data.")
    if not findings:
        findings.append("No additional table-specific structural risk was identified beyond the database-wide recommendations.")
    for finding in findings:
        add_bullet(doc, finding)

doc.add_page_break()
add_heading(doc, "Concise Final Risk and Recommendation Summary", 1)
for severity, recommendation in recommendations:
    add_bullet(doc, f"{severity}: {recommendation}")

doc.core_properties.title = "Comprehensive Database Summary - Setu_Odisha"
doc.core_properties.subject = "Database structure, data dictionary, relationships, objects and data quality"
doc.core_properties.keywords = "Setu_Odisha, SQL Server, database summary, data dictionary, data quality"
doc.core_properties.comments = "Generated from read-only database catalog and aggregate profiling."

doc.save(OUTPUT_PATH)
print(OUTPUT_PATH)
